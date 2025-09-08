import sys
import io
import unittest
from docx import Document

sys.path.append(".")
from app.agents.tools.word_document import WordDocumentInput, word_document_tool
from app.repositories.models.conversation import DocumentToolResultModel


class TestWordDocumentTool(unittest.TestCase):
    def test_word_document_creation(self):
        """Test that the word document tool creates a valid Word document."""
        # Test input
        content = "This is a test document.\nIt has multiple paragraphs.\n\nThis is another paragraph."
        title = "Test Document"
        
        # Create input object
        arg = WordDocumentInput(content=content, title=title)
        
        # Run the tool
        response = word_document_tool.run(
            tool_use_id="test_id",
            input=arg.model_dump(),
            model="claude-v3.5-sonnet-v2",
        )
        
        # Check response structure
        self.assertEqual(response["status"], "success")
        self.assertIsInstance(response["related_documents"], list)
        self.assertEqual(len(response["related_documents"]), 1)
        
        # Get the document from the response
        doc_result = response["related_documents"][0].content
        self.assertIsInstance(doc_result, DocumentToolResultModel)
        
        # Check document properties
        self.assertEqual(doc_result.format, "docx")
        self.assertEqual(doc_result.name, "Test_Document.docx")
        self.assertIsInstance(doc_result.document, bytes)
        
        # Verify the document content by loading it back
        doc_bytes = doc_result.document
        doc_stream = io.BytesIO(doc_bytes)
        doc = Document(doc_stream)
        
        # Check that the document has the expected content
        self.assertEqual(doc.paragraphs[0].text, title)  # First paragraph should be the title
        self.assertTrue(any("This is a test document" in p.text for p in doc.paragraphs))
        self.assertTrue(any("It has multiple paragraphs" in p.text for p in doc.paragraphs))
        self.assertTrue(any("This is another paragraph" in p.text for p in doc.paragraphs))

    def test_word_document_with_default_title(self):
        """Test that the word document tool works with default title."""
        # Test with only content, using default title
        content = "Simple content for testing default title."
        
        # Create input object with only content
        arg = WordDocumentInput(content=content)
        
        # Run the tool
        response = word_document_tool.run(
            tool_use_id="test_id",
            input=arg.model_dump(),
            model="claude-v3.5-sonnet-v2",
        )
        
        # Check response
        self.assertEqual(response["status"], "success")
        
        # Get the document from the response
        doc_result = response["related_documents"][0].content
        
        # Check document has default title in filename
        self.assertEqual(doc_result.name, "Generated_Document.docx")


if __name__ == "__main__":
    unittest.main()